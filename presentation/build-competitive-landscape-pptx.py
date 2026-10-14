# -*- coding: utf-8 -*-
"""One-page competitive-landscape slide for the QuantSim deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

FONT = "Segoe UI"
INK   = RGBColor(0x12, 0x17, 0x1C)
INK2  = RGBColor(0x3F, 0x4A, 0x55)
INK3  = RGBColor(0x6F, 0x7B, 0x87)
RULE  = RGBColor(0xC7, 0xD0, 0xD7)
RULE2 = RGBColor(0xE4, 0xE9, 0xEC)
YES   = RGBColor(0x0A, 0x6E, 0x4E)
PART  = RGBColor(0xC0, 0x7A, 0x06)
NO    = RGBColor(0xC0, 0x27, 0x2D)
QS    = RGBColor(0x0B, 0x6F, 0x5B)
QSBG  = RGBColor(0xEC, 0xF6, 0xF3)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

COLS = [
    ("SimCorp One",        "Institutional front-to-back"),
    ("QuantConnect",       "Cloud algo backtesting"),
    ("QuantRocket",        "Python research stack"),
    ("AlgoTrader / Wyden", "Institutional execution"),
    ("QuantSim",           "Research & risk simulation"),
]

ROWS = [
    ("Synthetic multi-model market simulation",
     "Monte Carlo, jump-diffusion, regime-switching path generation",
     ["P", "P", "N", "N", "Y"]),
    ("Deliberate stress & scenario testing",
     "Operator-defined shocks as a first-class workflow",
     ["Y", "N", "N", "P", "Y"]),
    ("Topology-based structural analysis",
     "Persistent homology on correlation structure",
     ["N", "N", "N", "N", "Y"]),
    ("Explainable AI on risk drivers",
     "Per-feature attribution behind every model output",
     ["N", "N", "N", "N", "Y"]),
    ("Composite risk score across lenses",
     "Stress, topology and ML reconciled into one verdict",
     ["P", "N", "N", "P", "Y"]),
    ("Strategy evaluation environment",
     "Run a strategy over generated or historical paths",
     ["P", "Y", "Y", "Y", "Y"]),
    ("Auditable, reproducible runs",
     "Versioned inputs, role-based access, traceable outputs",
     ["Y", "P", "N", "Y", "Y"]),
    ("Usable without institutional onboarding",
     "No enterprise sales cycle or per-module licensing",
     ["N", "Y", "Y", "N", "Y"]),
]

PRICING = [
    ("Enterprise licence, quote only",
     "Modular per-function and per-user subscription. Multi-year contracts. No public price list."),
    ("Free tier + seats",
     "Researcher ~$60/mo. Team from ~$24/user/mo. Trading Firm ~$48/user/mo. Institution ~$96/user/mo, billed annually."),
    ("Free tier + subscription",
     "Free tier with 2007-2011 US history. Paid from ~$20/mo. Non-Professional and Professional licences."),
    ("Enterprise licence, quote only",
     "Priced by users, markets traded and deployment model. Top of market. No public price list."),
    ("Single research environment",
     "One platform, no per-module licensing and no market-data resale. Proof of concept for research use."),
]

MARK = {"Y": (YES, u"\u2713"), "P": (PART, u"\u2013"), "N": (NO, u"\u2715")}

# ---------------------------------------------------------------- geometry
SW, SH = 13.333, 7.5
L, R = 0.42, 0.42
FEAT_W = 3.05
COL_W = (SW - L - R - FEAT_W) / 5.0
TABLE_TOP = 1.62
HEAD_H = 0.62
ROW_H = 0.44
BODY_TOP = TABLE_TOP + HEAD_H
BODY_BOT = BODY_TOP + ROW_H * len(ROWS)
PRICE_TOP = BODY_BOT + 0.14
PRICE_H = 1.14

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
slide = prs.slides.add_slide(prs.slide_layouts[6])
sh = slide.shapes


def col_x(i):
    return L + FEAT_W + COL_W * i


def rect(x, y, w, h, fill=None):
    s = sh.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.line.fill.background()
    s.shadow.inherit = False
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    s.text_frame.text = ""
    return s


def line(x, y, w, color=RULE2, pt=1.0):
    s = sh.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(pt))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def text(x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0):
    """runs: list of (string, size_pt, bold, color, space_before_pt) -> one paragraph each."""
    tb = sh.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for idx, spec in enumerate(runs):
        s, size, bold, color = spec[:4]
        before = spec[4] if len(spec) > 4 else 0
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if before:
            p.space_before = Pt(before)
        r = p.add_run()
        r.text = s
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def mark(cx, cy, kind, d=0.32):
    color, glyph = MARK[kind]
    s = sh.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = glyph
    r.font.name = FONT
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = WHITE


# ---------------------------------------------------------------- canvas
rect(0, 0, SW, SH, WHITE)

# QuantSim column band
rect(col_x(4), TABLE_TOP, COL_W, PRICE_TOP + PRICE_H - TABLE_TOP, QSBG)

# ---------------------------------------------------------------- header
text(L, 0.30, 9.0, 0.52, [("Competitive Landscape", 30, True, INK)], anchor=MSO_ANCHOR.TOP)
text(SW - R - 0.6, 0.32, 0.6, 0.4, [("4", 22, True, INK3)], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.TOP)
text(L, 0.86, 11.6, 0.62,
     [("Existing platforms are capable, but each is optimised around a different objective. The gap is the "
       "combination of synthetic multi-model simulation, deliberate stress testing, topology-based structural "
       "analysis and explainable AI inside one auditable research environment.", 12.5, False, INK2)],
     anchor=MSO_ANCHOR.TOP, spacing=1.15)

# ---------------------------------------------------------------- column heads
text(L, TABLE_TOP + 0.10, FEAT_W - 0.14, 0.40,
     [("CORE CAPABILITY", 10, True, INK3)], anchor=MSO_ANCHOR.BOTTOM)
for i, (name, kind) in enumerate(COLS):
    c = QS if i == 4 else INK
    ck = QS if i == 4 else INK3
    text(col_x(i) + 0.05, TABLE_TOP + 0.04, COL_W - 0.10, HEAD_H - 0.10,
         [(name, 13, True, c), (kind, 9, False, ck, 1)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM, spacing=1.0)

line(L, BODY_TOP - 0.02, SW - L - R, RULE, 1.6)

# ---------------------------------------------------------------- body rows
for r_i, (title, sub, marks) in enumerate(ROWS):
    y = BODY_TOP + ROW_H * r_i
    text(L, y + 0.03, FEAT_W - 0.16, ROW_H - 0.06,
         [(title, 11.5, True, INK), (sub, 9, False, INK3, 1)], spacing=1.05)
    for c_i, m in enumerate(marks):
        mark(col_x(c_i) + COL_W / 2, y + ROW_H / 2, m)
    if r_i < len(ROWS) - 1:
        line(L, y + ROW_H, SW - L - R, RULE2, 0.9)

# ---------------------------------------------------------------- pricing band
line(L, PRICE_TOP - 0.02, SW - L - R, RULE, 1.6)
text(L, PRICE_TOP + 0.06, FEAT_W - 0.16, 0.40, [("PRICING MODEL", 10, True, INK3)], anchor=MSO_ANCHOR.TOP)
for i, (headline, detail) in enumerate(PRICING):
    c = QS if i == 4 else INK
    text(col_x(i) + 0.07, PRICE_TOP + 0.06, COL_W - 0.14, PRICE_H - 0.10,
         [(headline, 10, True, c), (detail, 8.5, False, INK2, 2)],
         anchor=MSO_ANCHOR.TOP, spacing=1.1)

# ---------------------------------------------------------------- legend
ly = PRICE_TOP + PRICE_H + 0.20
lx = L
for kind, label in (("Y", "Core product capability"), ("P", "Partial or user-built"), ("N", "Not offered")):
    mark(lx + 0.11, ly, kind, d=0.22)
    text(lx + 0.28, ly - 0.11, 1.9, 0.22, [(label, 9, False, INK3)], anchor=MSO_ANCHOR.MIDDLE)
    lx += 2.15
text(SW - R - 5.6, ly - 0.20, 5.6, 0.42,
     [("QuantSim is a research and experimentation environment, not a replacement for institutional "
       "trading infrastructure. Pricing as published by each vendor, September 2026.", 8.5, False, INK3)],
     align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE, spacing=1.1)

out = r"C:\Users\YM\Desktop\QuantSim\QuantSim\presentation\QuantSim-Competitive-Landscape.pptx"
prs.save(out)
print("saved", out)
