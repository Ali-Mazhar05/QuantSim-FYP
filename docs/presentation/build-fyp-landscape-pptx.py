# -*- coding: utf-8 -*-
"""One-page prior-FYP comparison slide, styled to match the competitive-landscape slide."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

FONT = "Segoe UI"
TITLE = RGBColor(0x0B, 0x53, 0x34)   # dark forest green, as on the competitor slide
INK   = RGBColor(0x12, 0x17, 0x1C)
INK3  = RGBColor(0x6F, 0x7B, 0x87)
RULE  = RGBColor(0xC7, 0xD0, 0xD7)
RULE2 = RGBColor(0xE8, 0xEC, 0xEF)
YES   = RGBColor(0x0A, 0x6E, 0x4E)
PART  = RGBColor(0xC0, 0x7A, 0x06)
NO    = RGBColor(0xC0, 0x27, 0x2D)
QS    = RGBColor(0x0A, 0x6E, 0x4E)
QSBG  = RGBColor(0xEC, 0xF7, 0xF2)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

# Prior FAST FYPs closest to QuantSim, oldest to newest.
COLS = [("TRADINI",    "F22-066"),
        ("Binance ML",  "S23-027"),
        ("NeuraSight",  "F23-068"),
        ("TradeX",      "S25-024"),
        ("CashFlow",    "F25-091"),
        ("Trade360",    "F25-342"),
        ("QuantSim",    "this FYP")]

# Y = offered, P = partial / adjacent, N = not offered
ROWS = [("Synthetic market simulation",  ["N","N","N","N","N","P","Y"]),
        ("Stress & scenario testing",    ["N","N","N","N","N","P","Y"]),
        ("Topological analysis",         ["N","N","P","N","N","N","Y"]),
        ("Explainable AI on risk",       ["N","N","P","N","P","N","Y"]),
        ("Composite risk score",         ["N","N","N","P","N","P","Y"]),
        ("Strategy evaluation",          ["Y","Y","N","P","P","Y","Y"]),
        ("Auditable, reproducible runs", ["N","N","N","N","N","N","Y"]),
        ("Research output, not trading", ["N","N","Y","P","Y","N","Y"])]

MARK = {"Y": (YES, u"\u2713"), "P": (PART, u"\u2013"), "N": (NO, u"\u2715")}

# ---------------------------------------------------------------- geometry
SW, SH = 13.333, 7.5
L, R = 0.45, 0.45
FEAT_W = 2.85
COL_W = (SW - L - R - FEAT_W) / len(COLS)
HEAD_TOP = 1.28
BODY_TOP = 1.96
ROW_H = 0.585
BODY_BOT = BODY_TOP + ROW_H * len(ROWS)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(SW), Inches(SH)
slide = prs.slides.add_slide(prs.slide_layouts[6])
sh = slide.shapes


def col_x(i):
    return L + FEAT_W + COL_W * i


def rect(x, y, w, h, fill):
    s = sh.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background(); s.shadow.inherit = False
    s.text_frame.text = ""
    return s


def line(x, y, w, color=RULE2, pt=1.0):
    s = sh.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Pt(pt))
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    return s


def text(x, y, w, h, runs, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE, spacing=1.0):
    tb = sh.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for idx, spec in enumerate(runs):
        s, size, bold, color = spec[:4]
        before = spec[4] if len(spec) > 4 else 0
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        if before:
            p.space_before = Pt(before)
        r = p.add_run()
        r.text = s
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = color
    return tb


def mark(cx, cy, kind, d=0.36, glyph_pt=17):
    color, glyph = MARK[kind]
    s = sh.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    s.fill.solid(); s.fill.fore_color.rgb = color
    s.line.fill.background(); s.shadow.inherit = False
    tf = s.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = glyph
    r.font.name = FONT; r.font.size = Pt(glyph_pt)
    r.font.bold = True; r.font.color.rgb = WHITE


# ---------------------------------------------------------------- canvas
rect(0, 0, SW, SH, WHITE)
qi = len(COLS) - 1
rect(col_x(qi), HEAD_TOP - 0.1, COL_W, BODY_BOT - HEAD_TOP + 0.18, QSBG)

text(L, 0.40, 10.0, 0.62, [("Prior FYP Landscape", 32, True, TITLE)], anchor=MSO_ANCHOR.TOP)

# ---------------------------------------------------------------- headers
text(L, HEAD_TOP + 0.10, FEAT_W - 0.16, 0.42,
     [("CAPABILITY", 10.5, True, INK3)], anchor=MSO_ANCHOR.BOTTOM)
for i, (name, code) in enumerate(COLS):
    c  = QS if i == qi else INK
    ck = QS if i == qi else INK3
    text(col_x(i) + 0.04, HEAD_TOP, COL_W - 0.08, 0.62,
         [(name, 13, True, c), (code, 8.5, False, ck, 2)],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.BOTTOM)

line(L, BODY_TOP - 0.04, SW - L - R, RULE, 1.2)
line(col_x(qi), BODY_TOP - 0.05, COL_W, QS, 2.0)          # emphasis under QuantSim

# ---------------------------------------------------------------- body
for r_i, (label, marks) in enumerate(ROWS):
    y = BODY_TOP + ROW_H * r_i
    text(L, y, FEAT_W - 0.18, ROW_H, [(label, 12.5, True, INK)], spacing=1.05)
    for c_i, m in enumerate(marks):
        mark(col_x(c_i) + COL_W / 2, y + ROW_H / 2, m)
    if r_i < len(ROWS) - 1:
        line(L, y + ROW_H, SW - L - R, RULE2, 0.9)

# ---------------------------------------------------------------- legend
ly = BODY_BOT + 0.36
lx = L
for kind, label in (("Y", "Available"), ("P", "Partial / adjacent"), ("N", "Not offered")):
    mark(lx + 0.10, ly, kind, d=0.21, glyph_pt=10)
    text(lx + 0.27, ly - 0.11, 2.1, 0.22, [(label, 9.5, False, INK3)])
    lx += 1.95

out = r"C:\Users\YM\Desktop\QuantSim\QuantSim\presentation\QuantSim-FYP-Landscape.pptx"
prs.save(out)
print("saved", out)
print("bottom of table %.2f in, legend %.2f in, slide %.1f in" % (BODY_BOT, ly, SH))
print("column width %.3f in" % COL_W)
